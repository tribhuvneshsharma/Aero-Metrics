import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def build_presentation():
    src_template = r"C:\Users\Dell\Downloads\SIH2026-IDEA-Presentation-Format.pptx"
    out_workspace = r"c:\Users\Dell\Videos\Aero-Metrics\SIH2026_PS56_Aero-Metrics.pptx"
    out_downloads = r"C:\Users\Dell\Downloads\SIH2026_PS56_Aero-Metrics.pptx"

    prs = pptx.Presentation(src_template)

    # Color Palette
    C_NAVY = RGBColor(15, 23, 42)       # #0F172A (Primary text / Titles)
    C_SLATE = RGBColor(51, 65, 85)      # #334155 (Body text)
    C_BLUE = RGBColor(2, 132, 199)      # #0284C7 (Section Headers)
    C_DARKBLUE = RGBColor(30, 58, 138)  # #1E3A8A (Accent)

    def format_title(shape, title_text):
        shape.text = title_text
        for p in shape.text_frame.paragraphs:
            p.font.name = "Calibri"
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = C_NAVY

    def populate_textbox(shape, sections):
        shape.left = Inches(0.8)
        shape.top = Inches(1.35)
        shape.width = Inches(11.8)
        shape.height = Inches(5.3)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        tf.margin_top = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.margin_bottom = Inches(0.05)
        tf.clear()

        first = True
        for sec_idx, (sec_title, items) in enumerate(sections):
            p_head = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p_head.space_before = Pt(4) if sec_idx == 0 else Pt(8)
            p_head.space_after = Pt(2)

            r_head = p_head.add_run()
            r_head.text = sec_title
            r_head.font.name = "Calibri"
            r_head.font.size = Pt(13.5)
            r_head.font.bold = True
            r_head.font.color.rgb = C_BLUE

            for it in items:
                p_item = tf.add_paragraph()
                p_item.space_after = Pt(2.5)
                if "::" in it:
                    bp, rest = it.split("::", 1)
                    r_bp = p_item.add_run()
                    r_bp.text = "• " + bp.strip() + ": "
                    r_bp.font.name = "Calibri"
                    r_bp.font.size = Pt(11.5)
                    r_bp.font.bold = True
                    r_bp.font.color.rgb = C_NAVY

                    r_txt = p_item.add_run()
                    r_txt.text = rest.strip()
                    r_txt.font.name = "Calibri"
                    r_txt.font.size = Pt(11.5)
                    r_txt.font.color.rgb = C_SLATE
                else:
                    r_txt = p_item.add_run()
                    r_txt.text = "• " + it.strip()
                    r_txt.font.name = "Calibri"
                    r_txt.font.size = Pt(11.5)
                    r_txt.font.color.rgb = C_SLATE

    # ==================== SLIDE 1: TITLE PAGE ====================
    s1 = prs.slides[0]
    for sh in s1.shapes:
        if sh.name == "Subtitle 3":
            sh.text = "Aero-Metrics (APIx)\nGoverned High-Frequency Airfare Price Index Platform"
            for p in sh.text_frame.paragraphs:
                p.font.name = "Calibri"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = C_NAVY
        elif sh.name == "TextBox 9":
            tf = sh.text_frame
            tf.clear()
            fields = [
                ("Problem Statement ID -", " PS 56"),
                ("Problem Statement Title -", " Real-time Airfare Price Index (APIx)"),
                ("Theme -", " Smart Automation / Governance / FinTech"),
                ("PS Category -", " Software"),
                ("Team ID -", " [Insert Your Team ID]"),
                ("Team Name -", " Team Aero-Metrics")
            ]
            for i, (k, v) in enumerate(fields):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                r1 = p.add_run()
                r1.text = k
                r1.font.name = "Calibri"
                r1.font.size = Pt(14)
                r1.font.bold = True
                r1.font.color.rgb = C_NAVY

                r2 = p.add_run()
                r2.text = v
                r2.font.name = "Calibri"
                r2.font.size = Pt(14)
                r2.font.color.rgb = C_SLATE

    # ==================== SLIDE 2: IDEA TITLE ====================
    s2 = prs.slides[1]
    for sh in s2.shapes:
        if sh.name == "Title 1":
            format_title(sh, "IDEA TITLE: Aero-Metrics (APIx) - Policy-Grade Consumer Airfare Index")
        elif sh.name == "TextBox 8":
            sections_s2 = [
                ("1. Proposed Solution (Idea / Prototype Overview)", [
                    "Core Measurement Platform:: A governed statistical system converting volatile online consumer airfares into transparent, high-frequency price indices for RBI and NSO.",
                    "Fixed Sampling Basket:: Samples 16 high-density domestic directional corridors across 5 advance purchase horizons (T+1, T+7, T+15, T+30, T+45) at fixed collection windows (09:00 & 18:00 IST).",
                    "Policy-Grade Deliverables:: Publishes daily, weekly, and monthly headline and route-level indices accompanied by explicit coverage metrics and data quality scores."
                ]),
                ("2. How It Addresses the Problem", [
                    "Eliminates Official CPI Lag:: Bridges the 30-45 day publication lag of traditional CPI airfare series with high-frequency daily/weekly indicators.",
                    "Captures Booking Horizon Dynamics:: Replaces single spot quotes with multi-horizon booking curves reflecting real consumer purchasing behaviour.",
                    "Mandatory Consumer Payable Fare:: Aggregates true final payable fares (base fare + airline surcharge + UDF/PSF + statutory taxes), eliminating teaser fare distortion."
                ]),
                ("3. Innovation and Uniqueness of the Solution", [
                    "Full Traceability & Auditability:: Every index data point is fully traceable to its collection run ID, raw quote evidence pointer, and cleaning rule.",
                    "Bounded 2-Day Imputation Hierarchy:: Missing route observations are carried forward for max 2 days; beyond that, weights re-normalise automatically with partial coverage flags.",
                    "Ethical Web Governance:: Operates strictly via permitted public feeds, rate-limiting budgets, and robots rules - zero CAPTCHA evasion or legal exposure."
                ])
            ]
            populate_textbox(sh, sections_s2)

    # ==================== SLIDE 3: TECHNICAL APPROACH ====================
    s3 = prs.slides[2]
    for sh in s3.shapes:
        if sh.name == "Title 1":
            format_title(sh, "TECHNICAL APPROACH: System Architecture & Implementation Methodology")
        elif sh.name == "TextBox 8":
            sections_s3 = [
                ("1. Technologies to be Used", [
                    "Backend & APIs:: Python 3.12, FastAPI (high-performance REST & OpenAPI/Swagger), SQLAlchemy 2.0, Pydantic v2 (strict typed contracts).",
                    "Data Pipeline & Analytics:: Polars & Pandas (vectorised cleaning, canonical flight deduplication, MAD/IQR analytical outlier detection).",
                    "Storage Layer:: PostgreSQL (relational index and weights store) + MinIO / S3 (immutable raw quote audit store).",
                    "Frontend Dashboard:: Next.js (React 18 / TypeScript), Tailwind CSS, Recharts (executive dashboard, route heatmaps, booking curves).",
                    "DevOps & Orchestration:: Docker Compose (reproducible containerization), APScheduler (timed runs), Pytest automated test suite."
                ]),
                ("2. Methodology & Implementation Process", [
                    "Route Horizon Median:: P(r,h,t) = median(valid canonical total fares) - robust median eliminates distortion from extreme last-seat surge fares.",
                    "Route Price Relative:: R(r,h,t) = P(r,h,t) / P(r,h,0) benchmarked against a frozen base period (0).",
                    "Route Index Aggregation:: RouteIndex(r,t) = 100 * sum [v(h) * R(r,h,t)] using booking horizon weights (T+1: 10%, T+7: 25%, T+15: 30%, T+30: 20%, T+45: 15%).",
                    "National Headline APIx:: APIx(t) = sum [w(r) * RouteIndex(r,t)] using DGCA passenger traffic volume weights (summing exactly to 1.0)."
                ])
            ]
            populate_textbox(sh, sections_s3)

    # ==================== SLIDE 4: FEASIBILITY AND VIABILITY ====================
    s4 = prs.slides[3]
    for sh in s4.shapes:
        if sh.name == "Title 1":
            format_title(sh, "FEASIBILITY AND VIABILITY: Operational Feasibility, Challenges & Mitigation")
        elif sh.name == "TextBox 8":
            sections_s4 = [
                ("1. Analysis of Feasibility", [
                    "Deterministic 30-Day Replay Engine:: Ships pre-bundled with 30 dated daily snapshots (data/replay-30d/), guaranteeing a 100% reliable offline demo for hackathon evaluators regardless of live network connectivity or airline website status.",
                    "Modular Monorepo Architecture:: Decoupled contracts, collectors, pipeline, index engine, and frontend allow rapid scaling from prototype to production.",
                    "Lightweight Infrastructure:: Containerized architecture runs efficiently on local laptops, edge servers, or low-cost cloud virtual machines."
                ]),
                ("2. Potential Challenges & Risks", [
                    "Source Access Blocks & Volatility:: Airline websites changing DOM structure, rate-limiting, or blocking scraper requests.",
                    "Surge Fares & Extreme Outliers:: Severe price spikes caused by last-minute premium seat sales distorting macroeconomic inflation readings.",
                    "Missing Data Windows:: Sold-out flights or network drops causing incomplete route observations."
                ]),
                ("3. Strategies for Overcoming Challenges", [
                    "Robust Median Aggregation:: Eliminates distortion from extreme single-seat surge fares while preserving genuine route-level demand signals.",
                    "Transparent Statistical Quality Score:: Quality = 0.45*coverage + 0.25*success + 0.20*completeness + 0.10*freshness.",
                    "Fail-Safe Imputation Hierarchy:: Carries forward prices for max 2 days; suppresses headline index if weighted coverage falls below 70%."
                ])
            ]
            populate_textbox(sh, sections_s4)

    # ==================== SLIDE 5: IMPACT AND BENEFITS ====================
    s5 = prs.slides[4]
    for sh in s5.shapes:
        if sh.name == "Title 1":
            format_title(sh, "IMPACT AND BENEFITS: Target Audience Impact & Multi-Stakeholder Value")
        elif sh.name == "TextBox 8":
            sections_s5 = [
                ("1. Potential Impact on Target Audience", [
                    "Reserve Bank of India (RBI / MPC):: Provides a high-frequency supplementary gauge for transportation service inflation to inform quarterly monetary policy decisions.",
                    "Ministry of Civil Aviation (MoCA) & DGCA:: Granular route heatmap radar detects anti-competitive predatory pricing, calamity-driven surge gouging, and regional disparities.",
                    "National Statistical Office (NSO):: Offers a standardized, COICOP-aligned data pipeline to modernize official CPI baskets with high-frequency online quotes.",
                    "Consumers & General Public:: Visual lead-time booking curves (T+1 to T+45) inform travelers of optimal booking windows, minimizing travel costs."
                ]),
                ("2. Social, Economic & Governance Benefits", [
                    "Quantitative Policy Evidence:: Replaces anecdotal news-driven fare complaints with auditable, empirical price index data.",
                    "Machine-Readable Open Data:: RESTful OpenAPI endpoints enable direct algorithmic integration into government and institutional analytical systems.",
                    "Market Transparency & Fair Competition:: Encourages transparent pricing disclosures and curbs opportunistic fare surges on critical domestic corridors."
                ])
            ]
            populate_textbox(sh, sections_s5)

    # ==================== SLIDE 6: RESEARCH AND REFERENCES ====================
    s6 = prs.slides[5]
    for sh in s6.shapes:
        if sh.name == "Title 1":
            format_title(sh, "RESEARCH AND REFERENCES: Official Benchmarks, Standards & Project Links")
        elif sh.name == "TextBox 8":
            sections_s6 = [
                ("1. Official Data & Benchmark Sources", [
                    "DGCA India Passenger Traffic Reports (2025–2026):: Directional passenger volumes utilized to calibrate the 16 domestic corridor basket weights (w(r)).",
                    "DGCA Monthly Tariff Monitoring Data:: Official historical tariff benchmarks used to validate directional index movement during 30-day backtesting."
                ]),
                ("2. Methodological Standards & Academic Literature", [
                    "IMF Consumer Price Index Manual (Concepts and Methods):: Elementary price aggregation, price relatives formulation, and missing observation treatment.",
                    "Eurostat Practical Guidelines on Web-Scraped Data:: International guidelines on mitigating web collection sample bias in official price measurement.",
                    "MoSPI Technical Guidance Notes:: Consumer Price Index (CPI) methodology and weighting schemes for Indian transport and communication sub-indices."
                ]),
                ("3. Project Repository & Implementation Evidence", [
                    "Official GitHub Monorepo:: https://github.com/tribhuvneshsharma/Aero-Metrics",
                    "System Documentation & Methodology:: Full architectural, methodology, and governance specifications in repository /docs directory."
                ])
            ]
            populate_textbox(sh, sections_s6)

    # ==================== DELETE SLIDE 7 (INSTRUCTIONS) ====================
    # As explicitly instructed on Slide 7: "Note - You can delete this slide when you upload... maximum slides limit up to six (6)"
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Slide 7 (instructions) deleted successfully. Final presentation has exactly 6 slides.")

    # Save to workspace
    prs.save(out_workspace)
    print(f"Presentation saved to: {out_workspace}")

    # Save to downloads if not locked
    try:
        prs.save(out_downloads)
        print(f"Presentation copy saved to: {out_downloads}")
    except PermissionError:
        print(f"[Note] Could not overwrite {out_downloads} because it is currently open in PowerPoint. Workspace copy is updated.")


if __name__ == "__main__":
    build_presentation()
